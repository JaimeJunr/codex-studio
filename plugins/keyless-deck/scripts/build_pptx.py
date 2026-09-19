#!/usr/bin/env python3
"""Monta um .pptx a partir de uma pasta de slides-imagem (PNG full-bleed).

Uso:
    python build_pptx.py <dir_do_projeto> [--out arquivo.pptx] [--size 16:9|4:3|9:16]

Convenção da pasta (gerada pelo fluxo do keyless-deck):
    <dir>/origin_image/slide_01.png, slide_02.png, ...   (obrigatório)
    <dir>/speech.md                                        (opcional; vira notas)

Cada imagem entra como um slide full-bleed (a imagem É o slide — o keyless-deck
gera páginas inteiras). Sem dependência de fonte/tema do PowerPoint.

Requer: python-pptx  ->  pip install python-pptx
"""
import argparse
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    sys.exit("Falta python-pptx. Rode: pip install python-pptx")

# Tamanhos em EMU (1 inch = 914400 EMU).
SIZES = {
    "16:9": (12192000, 6858000),   # 13.333 x 7.5 in
    "4:3": (9144000, 6858000),     # 10 x 7.5 in
    "9:16": (6858000, 12192000),   # 7.5 x 13.333 in (vertical)
}


def natural_key(p: Path):
    """Ordena slide_2 antes de slide_10."""
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", p.name)]


def parse_speech(speech_path: Path):
    """Extrai notas por slide de um speech.md com marcadores '## Slide N'."""
    if not speech_path.exists():
        return {}
    notes, current = {}, None
    for line in speech_path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^#+\s*slide\s*(\d+)", line.strip(), re.IGNORECASE)
        if m:
            current = int(m.group(1))
            notes[current] = []
        elif current is not None:
            notes[current].append(line)
    return {k: "\n".join(v).strip() for k, v in notes.items()}


def build(project_dir: Path, out: Path, size: str):
    img_dir = project_dir / "origin_image"
    if not img_dir.is_dir():
        sys.exit(f"Nao encontrei {img_dir}. Gere os slides em origin_image/ primeiro.")
    slides = sorted(
        [p for p in img_dir.iterdir() if p.suffix.lower() in (".png", ".jpg", ".jpeg")],
        key=natural_key,
    )
    if not slides:
        sys.exit(f"Nenhuma imagem em {img_dir}.")

    width, height = SIZES[size]
    prs = Presentation()
    prs.slide_width = Emu(width)
    prs.slide_height = Emu(height)
    blank = prs.slide_layouts[6]  # layout em branco
    notes = parse_speech(project_dir / "speech.md")

    for i, img in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img), 0, 0, width=Emu(width), height=Emu(height))
        if i in notes and notes[i]:
            slide.notes_slide.notes_text_frame.text = notes[i]

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"OK: {len(slides)} slides -> {out}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("project_dir", help="Diretorio do projeto do deck")
    ap.add_argument("--out", default=None, help="Arquivo .pptx de saida")
    ap.add_argument("--size", default="16:9", choices=list(SIZES))
    args = ap.parse_args()

    project_dir = Path(args.project_dir).expanduser().resolve()
    out = Path(args.out) if args.out else project_dir / f"{project_dir.name}.pptx"
    build(project_dir, out, args.size)


if __name__ == "__main__":
    main()
